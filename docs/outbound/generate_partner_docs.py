#!/usr/bin/env python3
"""Generate two Word docs for partner outbound (architecture + collaboration)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

OUT = Path(__file__).resolve().parent
INK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
ACCENT = RGBColor(0x0E, 0x74, 0x90)
MUTED = RGBColor(0x64, 0x74, 0x8B)


def set_run(run, *, size=11, bold=False, color=SLATE, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_title(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=4, line=1.0)
    r = p.add_run(text)
    set_run(r, size=22, bold=True, color=INK, font="Calibri")
    return p


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=12, line=1.15)
    r = p.add_run(text)
    set_run(r, size=11, bold=False, color=MUTED)
    return p


def add_meta_line(doc, label, value):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=1.15)
    r1 = p.add_run(f"{label}: ")
    set_run(r1, size=10, bold=True, color=INK)
    r2 = p.add_run(value)
    set_run(r2, size=10, bold=False, color=SLATE)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=12, line=1.0)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CBD5E1")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    before = 18 if level == 1 else 12
    after = 8 if level == 1 else 6
    set_paragraph_spacing(p, before=before, after=after, line=1.15)
    r = p.add_run(text)
    if level == 1:
        set_run(r, size=14, bold=True, color=INK)
    else:
        set_run(r, size=12, bold=True, color=ACCENT)
    return p


def add_body(doc, text, *, justify=True):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=8, line=1.2)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run(r, size=11, color=SLATE)
    return p


def add_rich_body(doc, parts, *, justify=True):
    """parts: list of (text, bold)"""
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=8, line=1.2)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, bold in parts:
        r = p.add_run(text)
        set_run(r, size=11, bold=bold, color=SLATE)
    return p


def add_bullet(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=4, line=1.15)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run(r1, size=11, bold=True, color=INK)
        r2 = p.add_run(text)
        set_run(r2, size=11, color=SLATE)
    else:
        r = p.add_run(text)
        set_run(r, size=11, color=SLATE)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=10, line=1.25)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    r = p.add_run(text)
    set_run(r, size=11, bold=False, color=INK)
    r.italic = True
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, before=2, after=2, line=1.1)
        r = p.add_run(h)
        set_run(r, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "0B1F33")
        shading.set(qn("w:val"), "clear")
        cell._tc.get_or_add_tcPr().append(shading)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, before=2, after=2, line=1.15)
            r = p.add_run(str(val))
            set_run(r, size=10, color=SLATE)
            if ri % 2 == 1:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F8FAFC")
                shading.set(qn("w:val"), "clear")
                cell._tc.get_or_add_tcPr().append(shading)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph()
    return table


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = SLATE

    # List Bullet font
    try:
        lb = doc.styles["List Bullet"]
        lb.font.name = "Calibri"
        lb.font.size = Pt(11)
    except KeyError:
        pass

    return doc


def build_collaboration():
    doc = setup_doc()

    add_title(doc, "Collaboration Spike Outline")
    add_subtitle(doc, "Expert Scale × Overtone — Technical Collaboration Proposal")

    add_meta_line(doc, "Prepared for", "Drew Harris, Founder — Expert Scale (Apex Replicant™)")
    add_meta_line(doc, "Prepared by", "Muhammad Atif — Overtone")
    add_meta_line(doc, "Date", "August 2026")
    add_meta_line(doc, "Status", "Proposal only — no commercial commitment")
    add_hr(doc)

    add_heading(doc, "1. Purpose", 1)
    add_body(
        doc,
        "On our discussion you confirmed that Protégé runs primarily through a web interface today, "
        "and that live meeting join is the missing capability. Overtone already joins Meet, Zoom, and Teams "
        "as a participant and answers from an indexed deck or session source.",
    )
    add_heading(doc, "Escalation model (as you described)", 2)
    add_bullet(doc, " refuse / out of scope", bold_prefix="Out of specialty —")
    add_bullet(
        doc,
        " dig knowledge base deeper, then escalate if still insufficient",
        bold_prefix="On-scope but thin —",
    )
    add_bullet(
        doc,
        " escalate serious asks and time commitments to the human expert",
        bold_prefix="Lead-gen / limited capacity —",
    )
    add_rich_body(
        doc,
        [
            ("This spike wires ", False),
            ("(A) your judgment ladder", True),
            (" into ", False),
            ("(B) our live join path", True),
            (" once, end-to-end.", False),
        ],
    )

    add_heading(doc, "2. Goal (2–3 weeks, light)", 1)
    add_quote(
        doc,
        "Bot joins a live meeting → answers only when grounded in session sources → "
        "otherwise abstains or flags “needs expert” the way your escalate loop does.",
    )
    add_rich_body(
        doc,
        [
            ("This is a ", False),
            ("collaboration proof", True),
            (
                ", not a product sale and not an investment discussion. "
                "Commercial shape is deferred until after one clean end-to-end run.",
                False,
            ),
        ],
    )

    add_heading(doc, "3. Scope", 1)
    add_heading(doc, "In scope", 2)
    add_table(
        doc,
        ["#", "Deliverable"],
        [
            [
                "1",
                "Meet join (Recall) + realtime voice + deck/session grounding (existing Overtone path)",
            ],
            [
                "2",
                "Escalation mapping stub: out-of-scope refuse / on-scope dig / escalate-if-thin → log needs expert",
            ],
            [
                "3",
                "Short write-up: what worked, latency notes, and what would plug into Apex / Protégé APIs",
            ],
        ],
        col_widths=[0.5, 5.5],
    )

    add_heading(doc, "Out of scope for this spike", 2)
    add_bullet(doc, "Full Apex / Protégé product integration or production SLA")
    add_bullet(doc, "Pricing, equity, exclusivity, or fundraising")
    add_bullet(doc, "Rebuilding your judgment layer from scratch")

    add_heading(doc, "4. Success criteria", 1)
    add_body(doc, "We call the spike successful if:")
    add_bullet(
        doc,
        " a grounded question receives a source-bound answer (deck / session).",
        bold_prefix="In a live Meet,",
    )
    add_bullet(
        doc,
        " question refuses or escalates — no invented answer.",
        bold_prefix="An out-of-scope or ungroundable",
    )
    add_bullet(
        doc,
        " are visible (log / flag: topic, reason, timestamp) for a human expert to treat as ground truth later.",
        bold_prefix="Escalations",
    )

    add_heading(doc, "5. Contribution model", 1)
    add_table(
        doc,
        ["Expert Scale", "Overtone"],
        [
            [
                "Escalation rules + when “needs expert” is correct",
                "Meet / Zoom / Teams join + presenter AV + voice relay",
            ],
            [
                "Optional sample specialty scope (e.g. GTM-only) for the test",
                "Indexed deck / session grounding + tool calls",
            ],
            [
                "One short review call after you try Launch",
                "Spike build + honest latency / failure notes",
            ],
        ],
        col_widths=[3.0, 3.0],
    )

    add_heading(doc, "6. Suggested next steps", 1)
    add_bullet(
        doc,
        " on the dashboard with your own Meet link (feedback welcome).",
        bold_prefix="Try Launch",
    )
    add_bullet(
        doc,
        " to align escalate-stub fields with your handoff model today.",
        bold_prefix="Lock a 30-minute working session",
    )
    add_bullet(
        doc,
        " decide commercial shape (integration / ongoing collab) — or stop cleanly.",
        bold_prefix="After one clean end-to-end run,",
    )

    add_hr(doc)
    add_heading(doc, "Contact", 2)
    add_body(doc, "Muhammad Atif — Overtone", justify=False)
    add_body(doc, "ranaatif1299@gmail.com", justify=False)
    add_body(doc, "Dashboard (Launch): https://overtone-dashboard-4idrhaffca-uc.a.run.app", justify=False)

    path = OUT / "Overtone_Collaboration_Spike.docx"
    doc.save(path)
    return path


def build_architecture():
    doc = setup_doc()

    add_title(doc, "Architecture Overview")
    add_subtitle(doc, "Overtone — Meeting Join & Live Grounding")

    add_meta_line(doc, "Audience", "Technical partnership discussion")
    add_meta_line(doc, "Prepared by", "Muhammad Atif — Overtone")
    add_meta_line(doc, "Date", "August 2026")
    add_meta_line(doc, "Status", "Active development — demo / architecture stage")
    add_meta_line(doc, "Classification", "Confidential — partner review")
    add_hr(doc)

    add_heading(doc, "1. Framing", 1)
    add_rich_body(
        doc,
        [
            ("Overtone is a ", False),
            ("live meeting layer", True),
            (
                ": a bot joins Google Meet, Zoom, or Teams as a participant, presents a deck, "
                "listens with realtime speech-to-speech, and answers from ",
                False,
            ),
            ("indexed session sources", True),
            (
                " — with a design goal to abstain or escalate when grounding is weak (not invent).",
                False,
            ),
        ],
    )
    add_rich_body(
        doc,
        [
            ("It is ", False),
            ("not", True),
            (
                " a finished “zero-hallucination guarantee” product. The honest framing is: "
                "traceable, source-bound outputs plus a hard stop when the source cannot support an answer.",
                False,
            ),
        ],
    )

    add_heading(doc, "2. System context", 1)
    add_body(
        doc,
        "Operator uses the Dashboard to upload decks and launch bots. The Backend (FastAPI) creates a "
        "Recall.ai bot that joins the meeting room and opens the Presenter UI as the bot camera. "
        "The Presenter connects to the Backend over WebSocket for audio and tools. The Backend bridges "
        "to realtime voice (Gemini Live preferred, OpenAI Realtime as fallback) and retrieves grounded "
        "slide text from pgvector, with artifacts in GCS and session state in Postgres.",
    )
    add_table(
        doc,
        ["Component", "Role"],
        [
            ["Dashboard", "Upload decks, configure agents, launch bots into meetings"],
            ["Recall.ai", "Meeting participant join; virtual camera / audio path"],
            ["Presenter UI", "Deck display as bot camera; mic capture; play spoken audio"],
            ["Backend relay", "Session state, tool execution, voice bridge"],
            ["Realtime voice", "Speech-to-speech + function / tool calling"],
            ["pgvector + Vision", "Per-page extract, embed, retrieve grounded slide text"],
            ["GCS + Postgres", "Artifacts and durable session catalog"],
        ],
        col_widths=[1.7, 4.3],
    )

    add_heading(doc, "3. Session lifecycle", 1)
    add_bullet(
        doc,
        "Upload PDF / PPTX → store → Vision extract per page → quality gate → embeddings → pgvector → status ready",
    )
    add_bullet(
        doc,
        "Launch with meeting URL + knowledge base → create session → Recall bot with output-media presenter URL",
    )
    add_bullet(doc, "Bot joins room → Presenter WebSocket → Live voice session")
    add_bullet(
        doc,
        "Participant audio → relay → model; model tool calls → navigate / get_slide_details / search → spoken answer from tool results only",
    )

    add_heading(doc, "4. Platform layers", 1)
    add_table(
        doc,
        ["Layer", "Function", "Current technology"],
        [
            ["Capture / join", "Enter call as participant; camera shows slides", "Recall.ai"],
            ["Presenter", "Deck UI, mic, bot audio playback", "React (output-media page)"],
            ["Realtime voice", "Speech-to-speech + tools", "Gemini Live / OpenAI Realtime"],
            ["Orchestration", "State, tools, mute, page position", "FastAPI relay + executor"],
            ["Grounding", "Per-slide retrieval", "Vision → embeddings → pgvector"],
            ["Persistence", "Files, catalog, session extras", "GCS + Postgres"],
        ],
        col_widths=[1.4, 2.3, 2.3],
    )

    add_heading(doc, "5. Grounding path", 1)
    add_body(
        doc,
        "User question → tool routing (navigate / get_slide_details / search) → indexed slide text + session page. "
        "Enough evidence → spoken answer. Insufficient evidence → abstain / clarify → escalate path (roadmap).",
    )
    add_heading(doc, "Design principles (escalate-don’t-guess)", 2)
    add_bullet(doc, " Prefer the currently visible slide for “explain this / what’s on this.”")
    add_bullet(doc, " Search the deck index before any external lookup.")
    add_bullet(doc, " If retrieval is empty or stub-thin → do not invent; clarify or state not in source.")
    add_bullet(
        doc,
        " Human escalate (expert as ground truth) is the intended next architectural step — not claimed as fully shipped today.",
    )

    add_heading(doc, "6. Indexing pipeline", 1)
    add_body(
        doc,
        "Upload → store (GCS / local) → render page images → Vision per page → title / body / searchable_content → "
        "quality gate (reject stub pages) → embeddings → pgvector → presentation status ready. "
        "Vision stubs such as title-only “Slide 4” must not be marked ready.",
    )

    add_heading(doc, "7. Realtime tools", 1)
    add_table(
        doc,
        ["Tool", "Role"],
        [
            ["navigate_to_slide", "Jump to a page; persist current_page"],
            ["get_slide_details", "Narrate current or specified page from index"],
            ["search_and_answer", "Hybrid retrieve over slides; stay or navigate"],
            ["mute_self / unmute_self", "Hard silence path"],
            ["leave_call", "Explicit exit only"],
            ["fetch_external_data", "Last resort; not for on-deck questions"],
        ],
        col_widths=[2.0, 4.0],
    )
    add_rich_body(
        doc,
        [
            ("Spoken answers are instructed to come from ", False),
            ("tool results", True),
            (" (slide_content / brief_content), not model memory.", False),
        ],
    )

    add_heading(doc, "8. Maturity", 1)
    add_table(
        doc,
        ["Solid enough to demonstrate", "Explicit next (roadmap)"],
        [
            [
                "Join Meet via Recall; present deck as bot camera; realtime voice; tool-grounded slide Q&A; index quality gate",
                "First-class escalate-to-human queue; stronger abstention at tool boundary; multi-instance WS fan-out; claim→source audit trail",
            ]
        ],
        col_widths=[3.0, 3.0],
    )

    add_heading(doc, "9. Fit with a digital protégé product", 1)
    add_table(
        doc,
        ["Protégé concern", "Overtone role"],
        [
            ["Join Meet / Zoom as participant", "Meeting adapter (focus of current demos)"],
            ["Follow conversation without guessing", "Session + deck grounding + abstain"],
            ["Respond when addressed", "Realtime voice + tools"],
            ["Escalate when unsure", "Architectural fit — align with expert KB / human handoff"],
        ],
        col_widths=[2.5, 3.5],
    )
    add_rich_body(
        doc,
        [
            ("Overtone is best understood as a ", False),
            ("live meeting adapter + grounding shell", True),
            (", not a replacement for an expert-judgment knowledge system.", False),
        ],
    )

    add_hr(doc)
    add_body(
        doc,
        "This document describes current engineering architecture as of the date above. "
        "Capabilities labeled roadmap are intentional design direction, not production guarantees.",
        justify=False,
    )
    add_body(doc, "Contact: Muhammad Atif — ranaatif1299@gmail.com", justify=False)

    path = OUT / "Overtone_Architecture_Overview.docx"
    doc.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    a = build_architecture()
    c = build_collaboration()
    print(a)
    print(c)


if __name__ == "__main__":
    main()
